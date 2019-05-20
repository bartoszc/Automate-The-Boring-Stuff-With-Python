import openpyxl
from openpyxl.chart import BarChart, Reference

wb = openpyxl.Workbook()
chart = BarChart()
chart.title = 'First Series'

ws = wb.active
for i in range(1, 11):         # create some data in column A
    ws['A' + str(i)] = i
values = Reference(ws, min_col=1, min_row=1, max_col=1, max_row=10)
chart.add_data(values, titles_from_data=True)
ws.add_chart(chart, "E15")
wb.save('sampleChart.xlsx')
